#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador del deck "Sistema Ads + CRM para Clínicas" (omnia).
TEMPLATIZADO: editá el bloque CONFIG (marca, colores, logo, precios) y re-ejecutá:
    python build_deck.py
Modelo: omnia vende el sistema y factura todo · ómibu provee los ads (costo de omnia).
Salida: omnia-sistema-clinicas.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ============================ CONFIG (editar acá) ============================
BRAND      = "omnia"
SUBBRAND   = "Sistema Ads + CRM para Clínicas"
NICHOS     = "Odontología · Estética · Medicina estética"
ADS_PARTNER= "ómibu"                       # motor de ads (proveedor, uso interno)
LOGO_PATH  = ""                            # ruta a un PNG -> aparece en portada
PRIMARY    = RGBColor(0x0B, 0x1F, 0x3A)    # navy
ACCENT     = RGBColor(0x14, 0xB8, 0xA6)    # teal
ACCENT2    = RGBColor(0x25, 0x63, 0xEB)    # azul
WARN       = RGBColor(0x92, 0x40, 0x0E)
WARNBG     = RGBColor(0xFE, 0xF3, 0xC7)
INK        = RGBColor(0x0F, 0x17, 0x2A)
GRAY       = RGBColor(0x64, 0x74, 0x8B)
LIGHT      = RGBColor(0xF1, 0xF5, 0xF9)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
FONT       = "Calibri"
# Tarifas de gestión de ads (COSTO ómibu, €/mes) — reales del Excel 2026
ADS = {"Z":"250 €/mes","S":"450 €/mes","P":"625 €/mes","T":"a medida"}
OUT  = os.path.join(os.path.dirname(__file__), "omnia-sistema-clinicas.pptx")
TOTAL = 16
# ===========================================================================

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]; SW, SH = prs.slide_width, prs.slide_height

def slide(): return prs.slides.add_slide(BLANK)
def rect(s,l,t,w,h,fill,line=None,rounded=False):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(l),Inches(t),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(1)
    shp.shadow.inherit=False; return shp
def text(s,l,t,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,sp=4):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    if isinstance(runs[0],tuple): runs=[runs]
    for i,para in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(sp)
        for (txt,size,bold,color) in para:
            r=p.add_run(); r.text=txt; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name=FONT
    return tb
def footer(s,n,title="",internal=False):
    rect(s,0.55,6.92,12.2,0.012,LIGHT)
    tag = "USO INTERNO · " if internal else ""
    text(s,0.55,6.96,9,0.4,[[(f"{BRAND}  ·  {tag}{title}",9,internal,WARN if internal else GRAY)]])
    text(s,11.0,6.96,1.78,0.4,[[(f"{n:02d} / {TOTAL}",9,False,GRAY)]],align=PP_ALIGN.RIGHT)
def head(s,kicker,title,sub):
    rect(s,0.55,0.6,0.32,0.32,ACCENT,rounded=True)
    text(s,1.0,0.5,11,0.5,[[(kicker.upper(),11,True,ACCENT)]])
    text(s,0.55,0.95,12.2,0.9,[[(title,29,True,PRIMARY)]])
    if sub: text(s,0.55,1.72,12.2,0.6,[[(sub,15,False,GRAY)]])
def chain(s,items,top=3.2,h=1.5):
    n=len(items); gap=0.2; total=12.2; w=(total-gap*(n-1))/n
    for i,(t,d) in enumerate(items):
        l=0.55+i*(w+gap); rect(s,l,top,w,h,LIGHT,rounded=True)
        text(s,l+0.1,top+0.14,w-0.2,h-0.26,[[(t,12.5,True,PRIMARY)],[(d,10,False,GRAY)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        if i<n-1: text(s,l+w-0.02,top,gap+0.04,h,[[("›",18,True,ACCENT)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
def cards(s,items,top=2.6,h=3.2,cols=3):
    gap=0.3; total=12.2; w=(total-gap*(cols-1))/cols
    for i,(title,body,accent) in enumerate(items):
        row,col=divmod(i,cols); l=0.55+col*(w+gap); t=top+row*(h+0.3)
        rect(s,l,t,w,h,LIGHT,rounded=True); rect(s,l,t,w,0.13,accent)
        text(s,l+0.2,t+0.3,w-0.4,0.7,[[(title,14,True,PRIMARY)]])
        text(s,l+0.2,t+1.0,w-0.4,h-1.15,[[("•  "+b,11,False,INK)] for b in body],sp=5)

# 1 Portada
s=slide(); rect(s,0,0,SW.inches,SH.inches,PRIMARY); rect(s,0,0,0.25,SH.inches,ACCENT)
text(s,1.0,2.0,11.5,0.5,[[("GUÍA PARA EQUIPO Y CLIENTE",13,True,ACCENT)]])
text(s,1.0,2.5,11.5,1.2,[[(BRAND,54,True,WHITE)]])
text(s,1.0,3.7,11.5,0.9,[[(SUBBRAND,22,False,WHITE)]])
text(s,1.0,4.6,11.5,0.5,[[(NICHOS,15,False,ACCENT)]])
if LOGO_PATH and os.path.exists(LOGO_PATH): s.shapes.add_picture(LOGO_PATH,Inches(10.6),Inches(0.6),height=Inches(1.0))
text(s,1.0,6.3,11.5,0.5,[[("Convertimos anuncios en pacientes —  y pacientes en reseñas de 5★.",13,False,WHITE)]])

# 2 Problema
s=slide(); head(s,"El problema","El paciente escribe… y se pierde","Lo que hoy le pasa a casi toda clínica que invierte en publicidad.")
cards(s,[("Nadie contesta a tiempo",["El lead escribe a las 22:00","Le responden al otro día","Ya agendó en otra clínica"],ACCENT2),
("Todo disperso",["WhatsApp, Instagram, Facebook","Mensajes que se traspapelan","No se sabe de qué anuncio vino"],ACCENT2),
("Agenda con huecos",["Citas sin recordatorio","Ausencias y cancelaciones","Sin seguimiento al que no agendó"],ACCENT2),
("Reputación al azar",["Pocas reseñas en Google","Las malas sí se publican","Sin pedirlas de forma sistemática"],ACCENT2)],top=2.5,h=1.9,cols=2)
footer(s,2,"El problema")

# 3 Solución
s=slide(); head(s,"La solución","No es una herramienta. Es un SISTEMA.","Anuncios que traen pacientes + un CRM que los atiende y fideliza —  todo conectado.")
cards(s,[("1 · Captar",["Ads en Google/Meta","Atiende WhatsApp/IG/FB 24/7"],ACCENT),
("2 · Convertir",["IA reserva la cita","Mueve el pipeline solo"],ACCENT),
("3 · Fidelizar",["Recordatorios automáticos","Reseñas Google de 5★"],ACCENT)],top=2.7,h=2.3)
text(s,0.55,5.5,12.2,0.8,[[("→ Ads + CRM en un solo sistema. Una sola factura, un solo responsable: ",14,True,PRIMARY),(BRAND,14,True,ACCENT),(".",14,True,PRIMARY)]])
footer(s,3,"La solución")

# 4 Cómo funciona
s=slide(); head(s,"Cómo funciona","El recorrido del paciente, en piloto automático","De un clic en el anuncio a una reseña de 5★ —  sin que nadie mueva un dedo.")
chain(s,[("Anuncio","Google / Meta"),("Mensaje","WA/IG/FB"),("Sofía (IA)","atiende y califica"),("Agenda","reserva la cita"),("Recordatorios","24h antes"),("Reseña","Google 5★")],top=3.0,h=1.7)
text(s,0.55,5.4,12.2,0.9,[[("Resultado: el lead deja sus datos → la IA lo atiende → agenda solo → asiste → deja reseña → vuelve.",14,True,ACCENT)]])
footer(s,4,"Cómo funciona")

# 5 Sofía
s=slide(); head(s,"El asistente IA","Sofía: la recepcionista que no duerme","Atiende cada mensaje al instante, en el tono de la clínica, y reserva citas sola.")
cards(s,[("Atiende 24/7",["Responde en segundos","Sin sueldo ni descansos","Mismo tono de la marca"],ACCENT),
("Reserva citas",["Ofrece día y hora libres","Agenda en el calendario","Pide datos con RGPD"],ACCENT),
("Sabe cuándo parar",["Deriva a humano si hace falta","Comprobantes, dudas médicas","Nunca inventa precios"],ACCENT)],top=2.7,h=2.6)
footer(s,5,"Asistente IA")

# 6 Captación
s=slide(); head(s,"Captación","Todos los canales en una bandeja —  y sabés de dónde vino cada lead","")
cards(s,[("Multicanal",["WhatsApp, Instagram, Facebook","Una conversación por paciente","Nada se traspapela"],ACCENT2),
("Atribución de ads",["Guarda UTM y el anuncio (Ad)","Sabés qué campaña trae pacientes","Mides retorno por fuente"],ACCENT2),
("Canal detectado",["Etiqueta canal-wa/ig/fb","Pide el teléfono solo si hace falta","Listo para remarketing"],ACCENT2)],top=2.7,h=2.6)
footer(s,6,"Captación multicanal")

# 7 Pipeline
s=slide(); head(s,"Pipeline de ventas","Ves en qué punto está cada paciente","Cada oportunidad avanza por etapas claras —  como una operación con fases.")
chain(s,[("Nuevo Lead",""),("En conversación",""),("Pre-reserva",""),("Cita Reservada",""),("Cita Confirmada",""),("Servicio",""),("No reservó","")],top=3.2,h=1.4)
text(s,0.55,5.2,12.2,1.0,[[("Los robots mueven al paciente de etapa automáticamente. Vos ves el tablero y sabés qué falta para cada cierre.",14,False,GRAY)]])
footer(s,7,"Pipeline de ventas")

# 8 Automatizaciones
s=slide(); head(s,"Automatizaciones","9 robots que trabajan solos","Configurados y listos. Cada uno hace una tarea sin que nadie la recuerde.")
cards(s,[("Bienvenida + IA",["Activa a Sofía","Crea la oportunidad"],ACCENT),("Cita reservada",["Confirma por WhatsApp","Mueve el pipeline"],ACCENT),
("Recordatorio 24h",["Reduce ausencias","Mensaje automático"],ACCENT),("Remarketing",["Recupera al que no agendó","2 toques en 48h"],ACCENT),
("Señal / depósito",["Valida el pago","Confirma la cita"],ACCENT),("Reputación Google",["Pide reseña tras el servicio","Filtra 4-5★ → Google"],ACCENT)],top=2.5,h=1.85,cols=3)
footer(s,8,"Automatizaciones")

# 9 Reputación
s=slide(); head(s,"Reputación Google","Más reseñas de 5★ —  y las malas, en privado","El diferenciador: pide valoración tras cada servicio y filtra antes de Google.")
rect(s,0.55,2.7,5.9,2.6,LIGHT,rounded=True); rect(s,0.55,2.7,5.9,0.14,ACCENT)
text(s,0.8,3.0,5.4,2.2,[[("⭐ 4–5 estrellas",16,True,PRIMARY)],[("→ va directo a Google",12.5,False,INK)],[("Más reputación pública, más pacientes nuevos.",12,False,GRAY)]],sp=8)
rect(s,6.85,2.7,5.9,2.6,LIGHT,rounded=True); rect(s,6.85,2.7,5.9,0.14,ACCENT2)
text(s,7.1,3.0,5.4,2.2,[[("⭐ 1–3 estrellas",16,True,PRIMARY)],[("→ formulario privado (NO a Google)",12.5,False,INK)],[("Recuperás al paciente y proteges la marca.",12,False,GRAY)]],sp=8)
text(s,0.55,5.6,12.2,0.7,[[("También reactiva tu base de datos antigua para traer pacientes que ya tenías.",13.5,True,ACCENT)]])
footer(s,9,"Reputación Google")

# 10 Por nicho
s=slide(); head(s,"Adaptable por nicho","El mismo sistema, ajustado a cada clínica","Cambia el catálogo, el tono y si se pide señal. La estrategia es la misma.")
cards(s,[("Odontología",["Limpieza, ortodoncia, blanqueamiento","Tono cercano y claro","Señal: normalmente no"],ACCENT2),
("Estética / Belleza",["Faciales, micropigmentación","Tono cálido","Señal: habitual (sí)"],ACCENT),
("Medicina estética",["Ácido, bótox, tratamientos","Tono profesional","Dudas clínicas → a humano"],ACCENT2)],top=2.7,h=2.7)
footer(s,10,"Por nicho")

# 11 Templatizado
s=slide(); head(s,"Instalación","Templatizado: solo cambian 3–4 cosas","El sistema completo ya está montado. Personalizar un cliente toma horas, no semanas.")
cards(s,[("Se personaliza",["Logo y colores","Datos del negocio","Servicios y precios","¿Pide señal? sí/no"],ACCENT),
("Ya viene hecho",["Bot, pipeline y etapas","Los 9 workflows","Reputación Google","Plantillas de mensajes"],ACCENT2),
("Cómo se llena",["Custom Values (1 vez)","Todo lo lee el sistema","Mismo snapshot para todos","Onboarding por formulario"],ACCENT)],top=2.7,h=2.7)
footer(s,11,"Instalación templatizada")

# 12 Antes / Después
s=slide(); head(s,"El cambio","Antes vs. Después","Lo que ve el dueño de la clínica cuando enciende el sistema.")
rect(s,0.55,2.7,5.9,3.4,LIGHT,rounded=True); rect(s,0.55,2.7,5.9,0.14,GRAY)
text(s,0.85,3.0,5.4,3.0,[[("ANTES",15,True,GRAY)],[("—  Leads que escriben y nadie contesta",12.5,False,INK)],[("—  Agenda con huecos y ausencias",12.5,False,INK)],[("—  Mensajes perdidos entre apps",12.5,False,INK)],[("—  Pocas reseñas, y las malas públicas",12.5,False,INK)],[("—  Sin saber qué anuncio funciona",12.5,False,INK)]],sp=8)
rect(s,6.85,2.7,5.9,3.4,LIGHT,rounded=True); rect(s,6.85,2.7,5.9,0.14,ACCENT)
text(s,7.15,3.0,5.4,3.0,[[("DESPUÉS",15,True,ACCENT)],[("✓  Cada lead atendido en segundos, 24/7",12.5,False,INK)],[("✓  Agenda llena, con recordatorios",12.5,False,INK)],[("✓  Todo en una bandeja, nada se pierde",12.5,False,INK)],[("✓  Flujo constante de reseñas 5★",12.5,False,INK)],[("✓  Retorno medible por campaña",12.5,False,INK)]],sp=8)
footer(s,12,"Antes vs Después")

# 13 Planes (cliente) Z/S/P/T
s=slide(); head(s,"Planes","Cuatro niveles —  el anuncio trae, el sistema convierte","Cada plan = gestión de ads (Google/Meta) + el sistema CRM "+BRAND+" completo.")
planes=[("Captación · Z",["Ads: "+ADS["Z"],"1 canal (Google o Meta)","Campañas de leads","+ Sistema "+BRAND+" instalado"],ACCENT2),
("Crecimiento · S",["Ads: "+ADS["S"],"Más campañas / 2 canales","+ Analítica de leads","+ Sistema "+BRAND],ACCENT),
("Pro · P",["Ads: "+ADS["P"],"Gestión avanzada","+ Redistribución de presupuesto","+ Sistema "+BRAND],ACCENT2),
("Escala · T",["Ads: "+ADS["T"],"Multi-campaña / multi-sede","+ Reporting a medida","+ Sistema "+BRAND],ACCENT)]
cards(s,planes,top=2.5,h=2.85,cols=4)
rect(s,0.55,5.55,12.2,0.62,WARNBG,rounded=True)
text(s,0.75,5.62,11.8,0.5,[[("Inversión en anuncios aparte (al medio): mínimo 300 €/campaña.   ",11.5,True,WARN),("Precio al cliente y setup del sistema: a definir por "+BRAND+" (ver slide interna).",11.5,False,WARN)]])
footer(s,13,"Planes")

# 14 INTERNO — modelo
s=slide(); head(s,"Modelo de negocio","Cómo se arma y quién factura","Uso interno —  no mostrar al cliente.")
chain(s,[(ADS_PARTNER+" (Ads)","provee la gestión\nde campañas (COSTO)"),(BRAND+" (CRM)","añade el sistema\ny FACTURA todo"),("Cliente / Clínica","paga 1 sola\nfactura a "+BRAND)],top=3.0,h=1.8)
text(s,0.55,5.3,12.2,1.2,[[(BRAND+" compra los ads a "+ADS_PARTNER+" (tarifas del Excel = costo), le suma su producto propio —el sistema CRM (Sofía, pipeline, reputación)— y vende el paquete completo al cliente con una sola factura.",13.5,False,INK)]])
footer(s,14,"Modelo de negocio",internal=True)

# 15 INTERNO — costo / margen
s=slide(); head(s,"Costo y margen","Estructura por nivel (a completar)","Uso interno —  costos "+ADS_PARTNER+" reales; precio al cliente y margen los define "+BRAND+".")
cols=[("Nivel",2.6),("Ads (costo "+ADS_PARTNER+")",2.6),("Sistema "+BRAND,2.6),("Precio cliente",2.2),("Margen",2.2)]
x=0.55; t0=2.5
cx=x
for (h_,w_) in cols:
    rect(s,cx,t0,w_,0.5,PRIMARY); text(s,cx+0.08,t0,w_-0.16,0.5,[[(h_,10.5,True,WHITE)]],anchor=MSO_ANCHOR.MIDDLE); cx+=w_
rows=[("Captación · Z",ADS["Z"],"setup + mes","[definir]","[definir]"),
("Crecimiento · S",ADS["S"],"incluido","[definir]","[definir]"),
("Pro · P",ADS["P"],"incluido","[definir]","[definir]"),
("Escala · T",ADS["T"],"incluido","[definir]","[definir]")]
for r_i,row in enumerate(rows):
    ty=t0+0.5+r_i*0.55; cx=x; bg=WHITE if r_i%2==0 else LIGHT
    for (val,(h_,w_)) in zip(row,cols):
        rect(s,cx,ty,w_,0.55,bg)
        bold = (val.startswith("Captación") or val.startswith("Crecim") or val.startswith("Pro") or val.startswith("Escala"))
        col = WARN if val=="[definir]" else INK
        text(s,cx+0.08,ty,w_-0.16,0.55,[[(val,10.5,bold,col)]],anchor=MSO_ANCHOR.MIDDLE); cx+=w_
text(s,0.55,5.5,12.2,0.9,[[("Para cerrar la tabla necesito de "+BRAND+": precio del SISTEMA (setup + mensual) y si los ads van a costo o con markup.",12.5,True,WARN)]])
footer(s,15,"Costo / margen",internal=True)

# 16 Cierre
s=slide(); rect(s,0,0,SW.inches,SH.inches,PRIMARY); rect(s,0,0,0.25,SH.inches,ACCENT)
text(s,1.0,2.1,11.5,0.5,[[("EN RESUMEN",13,True,ACCENT)]])
text(s,1.0,2.6,11.5,1.6,[[("Un sistema que capta, atiende",30,True,WHITE)],[("y fideliza —  con una sola factura.",30,True,WHITE)]],sp=2)
text(s,1.0,4.4,11.5,1.4,[[("✓  Ads (Google/Meta) + CRM en un solo sistema",15,False,WHITE)],[("✓  En español, adaptable a odontología, estética y medicina",15,False,WHITE)],[("✓  Templatizado: nuevo cliente en horas",15,False,WHITE)]],sp=8)
text(s,1.0,6.3,11.5,0.5,[[(BRAND+"  ·  [web / contacto]",14,True,ACCENT)]])

prs.save(OUT); print("OK ->",OUT)
